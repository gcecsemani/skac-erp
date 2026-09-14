#!/usr/bin/env python3
"""Build a reviewable stock-on-hand UPDATE from Stock Master.xlsx.

  python3 scripts/generate_stock_master.py

Writes deploy/mysql/migrate_from_dump/21_stock_on_hand.sql

Columns used:
  Products / Packing  → product.name + packing
  TVM Shop            → branch 2 (TIRUVANNAMALAI)
  Avalurpet           → branch 1 (AVALURPET)
Godown is ignored. Blank TVM / Avalurpet cells are skipped for that branch.
Rows with both shop columns blank are skipped entirely.
"""
from __future__ import annotations

import re
from collections import defaultdict
from datetime import datetime, date
from decimal import Decimal, InvalidOperation
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
XLSX = ROOT / "Stock Master.xlsx"
PRODUCT_SQL = ROOT / "deploy" / "mysql" / "migrate_from_dump" / "08_product.sql"
BATCH_SQL = ROOT / "deploy" / "mysql" / "migrate_from_dump" / "10_batch.sql"
OUT = ROOT / "deploy" / "mysql" / "migrate_from_dump" / "21_stock_on_hand.sql"

AVL_BRANCH = 1
TVM_BRANCH = 2
ORG_ID = 1


ALIASES: dict[tuple[str, str], str] = {
    ("EM POWER", "250ML"): "EM POWER LIQ - 250 ML",
    ("HURRICAN PLUS", "100ML"): "HURRICANE PLUS - 100 ML",
    ("KANTROL PLUS", "100ML"): "KANTROL PLUSS - 100 ML",
    ("KICKER PLUS", "100ML"): "KICKER - 100 ML",
    ("OMITE", "250ML"): "OMITE 57 EC - 250MLS",
    ("VAJRA 19 19 19", "100ML"): "19:19:19 VAJRA - 100 ML",
    ("SAI POWER", "100ML"): "SAI POWER PLUS - 100 ML",
    ("PRETI SUPER", "600ML"): "PRETI EW - 600 ML",
}


def sql_str(value: object | None) -> str:
    if value is None:
        return "NULL"
    return "'" + str(value).replace("\\", "\\\\").replace("'", "\\'") + "'"


def sql_qty(value: Decimal) -> str:
    if value == value.to_integral():
        return str(int(value))
    return format(value, "f")


def as_qty(value: object) -> Decimal | None:
    if value is None:
        return None
    if isinstance(value, str) and not value.strip():
        return None
    try:
        q = Decimal(str(value).strip())
    except (InvalidOperation, ValueError, TypeError):
        return None
    return q


def norm_name(text: str) -> str:
    s = str(text or "").upper().replace("&", " AND ")
    s = re.sub(r"[^A-Z0-9]+", " ", s)
    return re.sub(r"\s+", " ", s).strip()


def compact_name(text: str) -> str:
    return re.sub(r"[^A-Z0-9]", "", norm_name(text))


def norm_pack(text: str) -> str:
    s = re.sub(r"\s+", "", str(text or "").upper())
    reps = (
        "MILLILITRES",
        "MILLILITERS",
        "MILLILITRE",
        "MILLILITER",
        "MLS",
        "LITRES",
        "LITERS",
        "LITRE",
        "LTRS",
        "LTR",
        "LITR",
        "LTRE",
        "KILOGRAMS",
        "KILOGRAM",
        "KGS",
        "GRAMS",
        "GRAM",
        "GMS",
    )
    units = {
        "MILLILITRES": "ML",
        "MILLILITERS": "ML",
        "MILLILITRE": "ML",
        "MILLILITER": "ML",
        "MLS": "ML",
        "LITRES": "L",
        "LITERS": "L",
        "LITRE": "L",
        "LTRS": "L",
        "LTR": "L",
        "LITR": "L",
        "LTRE": "L",
        "KILOGRAMS": "KG",
        "KILOGRAM": "KG",
        "KGS": "KG",
        "GRAMS": "G",
        "GRAM": "G",
        "GMS": "G",
    }
    for token in reps:
        if s.endswith(token):
            s = s[: -len(token)] + units[token]
            break
    if s.endswith("GM") and not s.endswith("KG"):
        s = s[:-2] + "G"
    if s.endswith("MML"):
        s = s[:-3] + "ML"
    if s.endswith("MO"):  # occasional packing typo for ML
        s = s[:-2] + "ML"
    return s


def parse_products() -> list[dict]:
    text = PRODUCT_SQL.read_text(encoding="utf-8")
    rows = []
    for line in text.splitlines():
        if not line.startswith("  ("):
            continue
        m = re.match(
            r"\s+\((\d+), 1, 'P\d+', (?:NULL|'[^']*'), '((?:\\'|[^'])*)',",
            line,
        )
        if not m:
            continue
        pid = int(m.group(1))
        name = m.group(2).replace("\\'", "'")
        pack_m = re.search(r'"packing":"([^"]*)"', line)
        pack = pack_m.group(1) if pack_m else ""
        end = re.search(r", (\d)\),?\s*$", line)
        deleted = bool(end and end.group(1) == "1")
        if " - " in name:
            item, name_pack = name.rsplit(" - ", 1)
        else:
            item, name_pack = name, pack
        rows.append(
            {
                "id": pid,
                "name": name,
                "item": item,
                "pack": pack or name_pack,
                "deleted": deleted,
            }
        )
    return rows


def parse_opening_batches() -> dict[int, int]:
    text = BATCH_SQL.read_text(encoding="utf-8")
    out: dict[int, int] = {}
    for m in re.finditer(r"\((\d+), 1, (\d+), 'LEGACY-OPENING'", text):
        out[int(m.group(2))] = int(m.group(1))
    return out


def prefer(cands: list[dict]) -> dict | None:
    live = [c for c in cands if not c["deleted"]] or cands
    uniq = {c["id"]: c for c in live}
    if len(uniq) == 1:
        return next(iter(uniq.values()))
    # newest id among live
    if not live:
        return None
    return max(live, key=lambda c: c["id"])


class Matcher:
    def __init__(self, products: list[dict]):
        self.products = products
        self.by_key: dict[tuple[str, str], list[dict]] = defaultdict(list)
        self.by_compact: dict[tuple[str, str], list[dict]] = defaultdict(list)
        self.by_pack: dict[str, list[dict]] = defaultdict(list)
        self.by_full: dict[str, dict] = {}
        for p in products:
            key = (norm_name(p["item"]), norm_pack(p["pack"]))
            self.by_key[key].append(p)
            self.by_compact[(compact_name(p["item"]), norm_pack(p["pack"]))].append(p)
            self.by_pack[norm_pack(p["pack"])].append(p)
            self.by_full[p["name"].upper()] = p

    def match(self, product: str, packing: str) -> tuple[dict | None, str]:
        n = norm_name(product)
        p = norm_pack(packing)
        alias = ALIASES.get((n, p))
        if alias:
            hit = self.by_full.get(alias.upper())
            if hit:
                return hit, "alias"
        c = compact_name(product)
        hits = self.by_key.get((n, p), [])
        if hits:
            return prefer(hits), "exact"
        hits = self.by_compact.get((c, p), [])
        if hits:
            return prefer(hits), "compact-name"
        sheet_tokens = [t for t in n.split() if t]
        scored: list[tuple[tuple, dict]] = []
        for row in self.by_pack.get(p, []):
            item_n = norm_name(row["item"])
            item_tokens = item_n.split()
            item_set = set(item_tokens)
            if not sheet_tokens:
                continue
            overlap = sum(1 for t in sheet_tokens if t in item_set)
            if overlap == 0:
                continue
            standalone = compact_name(row["item"]) in {"PLUS", "POW", "GR", "HD"}
            score = (overlap, 0 if not standalone else -5, len(item_n))
            if overlap == len(sheet_tokens) or (overlap >= 2 and overlap >= len(sheet_tokens) - 1):
                scored.append((score, row))
        if not scored:
            return None, "unmatched"
        scored.sort(key=lambda x: x[0], reverse=True)
        best_score = scored[0][0]
        best = [row for score, row in scored if score[:2] == best_score[:2]]
        live = [x for x in best if not x["deleted"]] or best
        ids = {x["id"] for x in live}
        if len(ids) == 1:
            return prefer(live), "packing+tokens"
        return None, "ambiguous:" + "/".join(sorted({x["name"] for x in live})[:4])


def fmt_expiry(value: object) -> str | None:
    if value is None:
        return None
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    text = str(value).strip()
    if not text:
        return None
    for fmt in ("%Y-%m-%d", "%d-%m-%Y", "%d/%m/%Y", "%Y/%m/%d"):
        try:
            return datetime.strptime(text[:10], fmt).date().isoformat()
        except ValueError:
            continue
    return None


def load_sheet() -> list[dict]:
    import sys

    sys.path.insert(0, "/tmp/pydeps")
    import openpyxl

    wb = openpyxl.load_workbook(XLSX, data_only=True)
    ws = wb["Sheet1"]
    rows = []
    for r in ws.iter_rows(min_row=2, max_row=ws.max_row, max_col=12, values_only=True):
        name = (str(r[1]).strip() if r[1] is not None else "")
        pack = (str(r[2]).strip() if r[2] is not None else "")
        if not name:
            continue
        rows.append(
            {
                "sno": r[0],
                "name": name,
                "pack": pack,
                "tvm": as_qty(r[3]),
                "avl": as_qty(r[5]),
                "expiry": fmt_expiry(r[7]),
                "company": r[9],
                "price": r[10],
            }
        )
    return rows


def main() -> None:
    products = parse_products()
    matcher = Matcher(products)
    sheet = load_sheet()
    print(f"Products {len(products)}  sheet rows {len(sheet)}")

    matched: list[dict] = []
    unmatched: list[dict] = []
    skipped = 0
    for row in sheet:
        if row["tvm"] is None and row["avl"] is None:
            skipped += 1
            continue
        hit, how = matcher.match(row["name"], row["pack"])
        if not hit:
            unmatched.append({**row, "how": how})
            continue
        matched.append({**row, "product": hit, "how": how, "avl_qty": row["avl"]})

    product_ids = sorted({m["product"]["id"] for m in matched})
    id_list = ", ".join(str(i) for i in product_ids)
    lines: list[str] = [
        "-- Physical stock on hand from Stock Master.xlsx (Sheet1)",
        "-- Review before running. Safe to re-run.",
        "--",
        "-- Branch map:",
        "--   Avalurpet column  →  branch 1  AVALURPET (AVL)",
        "--   TVM Shop column   →  branch 2  TIRUVANNAMALAI (TVM)",
        "--   Godown column     →  ignored",
        "--",
        "-- Skip rules: blank TVM / Avalurpet is skipped for that branch.",
        "-- Rows with both shop columns blank are skipped. 0 is treated as counted (stock set to 0).",
        "-- Counted qty is stored on a STOCK-ON-HAND batch looked up/created by product_id",
        "-- (does not depend on dump batch ids, which may not exist in a live DB).",
        "--",
        f"-- Sheet rows: {len(sheet)}  skipped (no qty): {skipped}",
        f"-- Matched with qty: {len(matched)}  unmatched: {len(unmatched)}",
        "--",
        "-- STEP 1 zeros dump opening stock (LEGACY-OPENING) on both branches.",
        "-- Comment out STEP 1 if you only want to overlay counted SKUs.",
        "--",
        "--   mysql -u root -p skac_new < deploy/mysql/migrate_from_dump/21_stock_on_hand.sql",
        "",
        "USE skac_new;",
        "SET NAMES utf8mb4;",
        "",
        "-- ---------------------------------------------------------------------------",
        "-- STEP 0) Opening batch per counted product (id comes from live batch table)",
        "-- ---------------------------------------------------------------------------",
        "INSERT INTO batch (organization_id, product_id, batch_no, expiry_date, purchase_price)",
        "SELECT p.organization_id, p.id, 'STOCK-ON-HAND', NULL, COALESCE(p.purchase_price, 0)",
        "FROM product p",
        f"WHERE p.organization_id = {ORG_ID} AND p.id IN ({id_list})",
        "  AND NOT EXISTS (",
        "    SELECT 1 FROM batch b",
        "    WHERE b.organization_id = p.organization_id",
        "      AND b.product_id = p.id",
        "      AND b.batch_no = 'STOCK-ON-HAND'",
        "  );",
        "",
        "-- ---------------------------------------------------------------------------",
        "-- STEP 1) Clear dump opening stock (LEGACY-OPENING) on both shops",
        "-- ---------------------------------------------------------------------------",
        "UPDATE stock s",
        "JOIN batch b ON b.id = s.batch_id AND b.batch_no = 'LEGACY-OPENING'",
        "SET s.quantity = 0, s.updated_at = NOW()",
        f"WHERE s.organization_id = {ORG_ID};",
        "",
        "-- ---------------------------------------------------------------------------",
        "-- STEP 2) Set counted quantities on STOCK-ON-HAND",
        "-- ---------------------------------------------------------------------------",
        "",
    ]

    avl_n = tvm_n = 0
    avl_sum = tvm_sum = Decimal("0")
    for row in sorted(matched, key=lambda x: (x["product"]["name"], x["product"]["id"])):
        p = row["product"]
        pid = p["id"]
        tvm = row["tvm"]
        avl_qty = row["avl_qty"]
        parts = [
            f"TVM={tvm if tvm is not None else '-'}",
            f"Avalurpet={avl_qty if avl_qty is not None else '-'}",
        ]
        lines.append(f"-- {p['name']}  (product_id={pid}, match={row['how']})")
        lines.append(f"--   sheet {row['name']} / {row['pack']}   " + "  ".join(parts))
        if avl_qty is not None:
            avl_n += 1
            avl_sum += avl_qty
            lines.append(
                "INSERT INTO stock (organization_id, branch_id, product_id, batch_id, quantity) "
                f"SELECT {ORG_ID}, {AVL_BRANCH}, {pid}, b.id, {sql_qty(avl_qty)} "
                "FROM batch b "
                f"WHERE b.organization_id = {ORG_ID} AND b.product_id = {pid} AND b.batch_no = 'STOCK-ON-HAND' "
                "AND NOT EXISTS ("
                f"SELECT 1 FROM stock s WHERE s.branch_id = {AVL_BRANCH} AND s.batch_id = b.id"
                ");"
            )
            lines.append(
                "UPDATE stock s "
                "JOIN batch b ON b.id = s.batch_id AND b.batch_no = 'STOCK-ON-HAND' "
                f"AND b.product_id = {pid} "
                f"SET s.quantity = {sql_qty(avl_qty)}, s.updated_at = NOW() "
                f"WHERE s.organization_id = {ORG_ID} AND s.branch_id = {AVL_BRANCH} AND s.product_id = {pid};"
            )
        if tvm is not None:
            tvm_n += 1
            tvm_sum += tvm
            lines.append(
                "INSERT INTO stock (organization_id, branch_id, product_id, batch_id, quantity) "
                f"SELECT {ORG_ID}, {TVM_BRANCH}, {pid}, b.id, {sql_qty(tvm)} "
                "FROM batch b "
                f"WHERE b.organization_id = {ORG_ID} AND b.product_id = {pid} AND b.batch_no = 'STOCK-ON-HAND' "
                "AND NOT EXISTS ("
                f"SELECT 1 FROM stock s WHERE s.branch_id = {TVM_BRANCH} AND s.batch_id = b.id"
                ");"
            )
            lines.append(
                "UPDATE stock s "
                "JOIN batch b ON b.id = s.batch_id AND b.batch_no = 'STOCK-ON-HAND' "
                f"AND b.product_id = {pid} "
                f"SET s.quantity = {sql_qty(tvm)}, s.updated_at = NOW() "
                f"WHERE s.organization_id = {ORG_ID} AND s.branch_id = {TVM_BRANCH} AND s.product_id = {pid};"
            )
        if row["expiry"]:
            lines.append(
                "UPDATE batch SET "
                f"expiry_date = {sql_str(row['expiry'])}, updated_at = NOW() "
                f"WHERE organization_id = {ORG_ID} AND product_id = {pid} AND batch_no = 'STOCK-ON-HAND' "
                f"AND (expiry_date IS NULL OR expiry_date <> {sql_str(row['expiry'])});"
            )
        lines.append("")

    lines += [
        "-- ---------------------------------------------------------------------------",
        "-- Unmatched sheet rows (not applied). Add aliases or create the packing in product.",
        "-- ---------------------------------------------------------------------------",
        "",
    ]
    for row in unmatched:
        t = row["tvm"] if row["tvm"] is not None else "-"
        a = row["avl"] if row["avl"] is not None else "-"
        lines.append(
            f"-- {row['name']} / {row['pack']}   TVM={t}  Avalurpet={a}   ({row['how']})"
        )

    lines += [
        "",
        f"-- Applied AVL rows: {avl_n}  qty {avl_sum}",
        f"-- Applied TVM rows: {tvm_n}  qty {tvm_sum}",
        f"-- Unmatched counted rows: {len(unmatched)}",
        "",
    ]
    OUT.write_text("\n".join(lines) + "\n", encoding="utf-8")
    print(f"Wrote {OUT}")
    print(f"matched {len(matched)} unmatched {len(unmatched)} skipped {skipped}")
    print(f"AVL lines {avl_n} qty {avl_sum}  TVM lines {tvm_n} qty {tvm_sum}")
    print("sample unmatched:")
    for row in unmatched[:12]:
        print(f"  {row['name']} / {row['pack']}  {row['how']}")


if __name__ == "__main__":
    main()
